
from docx import Document as WordDocument
from reportlab.pdfgen import canvas

import pytest

from src.ingestion.file_loader import load_document


def test_load_txt_document(tmp_path):
    path = tmp_path / "incident.txt"

    path.write_text(
        "Payment Processing Incident\n\n"
        "The payment service was unavailable for ten minutes. "
        "The team rolled back the deployment.",
        encoding="utf-8",
    )

    result = load_document(path)

    assert result["title"] == "incident"
    assert result["metadata"]["file_format"] == "txt"
    assert "payment service" in result["text"].lower()
    assert len(result["sentences"]) >= 2
    assert result["provenance"][0]["original_filename"] == (
        "incident.txt"
    )


def test_load_docx_document_and_table(tmp_path):
    path = tmp_path / "incident.docx"

    document = WordDocument()
    document.add_heading("Payment Incident", level=1)
    document.add_paragraph(
        "The payment service experienced an outage."
    )

    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Service"
    table.cell(0, 1).text = "Payment API"

    document.save(path)

    result = load_document(path)

    assert result["metadata"]["file_format"] == "docx"
    assert "payment service" in result["text"].lower()
    assert "payment api" in result["text"].lower()
    assert len(result["sentences"]) >= 2


def test_load_pdf_document(tmp_path):
    path = tmp_path / "incident.pdf"

    pdf = canvas.Canvas(str(path))
    pdf.drawString(
        72,
        720,
        "The payment service was unavailable for ten minutes.",
    )
    pdf.drawString(
        72,
        690,
        "The engineering team restored the service.",
    )
    pdf.save()

    result = load_document(path)

    assert result["metadata"]["file_format"] == "pdf"
    assert result["metadata"]["page_count"] == 1
    assert "payment service" in result["text"].lower()
    assert len(result["sentences"]) >= 2


def test_same_filename_has_stable_document_id(tmp_path):
    path = tmp_path / "notes.txt"

    path.write_text(
        "First version of a document.",
        encoding="utf-8",
    )
    first = load_document(path)

    path.write_text(
        "Updated version of the document.",
        encoding="utf-8",
    )
    second = load_document(path)

    assert first["document_id"] == second["document_id"]
    assert (
        first["metadata"]["content_sha256"]
        != second["metadata"]["content_sha256"]
    )


def test_unsupported_file_type_is_rejected(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text("name,value\nA,1", encoding="utf-8")

    with pytest.raises(
        ValueError,
        match="Unsupported file format",
    ):
        load_document(path)


def test_empty_text_file_is_rejected(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_text("", encoding="utf-8")

    with pytest.raises(
        ValueError,
        match="no extractable text",
    ):
        load_document(path)
